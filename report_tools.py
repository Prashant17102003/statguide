"""Build one report (list of blocks) and render it as Word, PDF or Excel."""
import datetime
import io
import math

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image as PILImage

import eda_tools as E
import pca_tools as PC
import regression_tools as RG
from advisor_tools import GOALS, normality_table, group_normality, overall, recommend
from stats_tools import (two_group_test, multi_group_test, interpret_anova, chi_square_test, interpret_chi,
                         paired_test, one_sample_test, friedman_test, interpret_paired, interpret_one_sample, interpret_friedman, fmt_p)

SECTIONS = ["Dataset overview", "Descriptive statistics", "Missing values", "Normality assessment",
            "Distributions (first variables)", "Correlation", "Test advisor", "PCA"]


def png(fig):
    b = io.BytesIO(); fig.savefig(b, format="png", dpi=110, bbox_inches="tight"); plt.close(fig)
    return b.getvalue()


def stitch(pngs):
    ims = [PILImage.open(io.BytesIO(p)).convert("RGB") for p in pngs]
    c = PILImage.new("RGB", (sum(i.width for i in ims), max(i.height for i in ims)), "white"); x = 0
    for i in ims:
        c.paste(i, (x, 0)); x += i.width
    b = io.BytesIO(); c.save(b, format="PNG")
    return b.getvalue()


def _cell(v):
    if isinstance(v, (float, np.floating)):
        return "" if np.isnan(v) else f"{v:.4g}"
    return str(v)


def fmt(df, max_rows=40):
    d = df.copy()
    if not isinstance(d.index, pd.RangeIndex):
        if d.index.name is None:
            d.index.name = "variable"
        d = d.reset_index()
    note = f"Showing the first {max_rows} of {len(d)} rows." if len(d) > max_rows else None
    d = d.head(max_rows)
    return [str(c) for c in d.columns], [[_cell(v) for v in row] for row in d.itertuples(index=False)], note


def build_report(df, name, cfg):
    B, inc = [], set(cfg.get("include", SECTIONS)); nc, cc = E.split_columns(df); num = [0]
    def h1(t):
        num[0] += 1; B.append(("h1", f"{num[0]}. {t}"))
    p = lambda t: B.append(("p", t)); bl = lambda L: B.append(("bullets", list(L)))
    tb = lambda d, cap: B.append(("table", d, cap)); im = lambda b, cap: B.append(("img", b, cap))
    B.append(("title", cfg.get("title") or "Data analysis report",
              f"Source: {name}  |  Generated: {datetime.date.today():%d %b %Y}" + (f"  |  Prepared by: {cfg['author']}" if cfg.get("author") else "")))
    prof = E.profile_columns(df)
    normal, nt = False, pd.DataFrame()
    if nc:
        nt = normality_table(df, 0.05)
        if not nt.empty:
            normal = bool(overall(nt, 80)[0])
    if "Dataset overview" in inc:
        h1("Dataset overview")
        p(f"The dataset has **{len(df)} rows** and **{df.shape[1]} columns** ({len(nc)} numeric, {len(cc)} categorical or binary) and "
          f"**{int(df.isna().sum().sum())} missing cells**.")
        tb(prof, "Column types and missing values")
    if "Descriptive statistics" in inc:
        h1("Descriptive statistics")
        if nc:
            d = df[nc].describe().T[["count", "mean", "std", "min", "25%", "50%", "75%", "max"]]
            d["skew"], d["kurtosis"] = df[nc].skew(), df[nc].kurt(); tb(d.round(4), "Numeric variables")
        rows = []
        for c in cc[:30]:
            vc = df[c].value_counts()
            if len(vc):
                rows.append({"variable": c, "categories": len(vc), "most common": str(vc.index[0]),
                             "share %": round(100 * vc.iloc[0] / vc.sum(), 1), "missing %": round(100 * df[c].isna().mean(), 1)})
        if rows:
            tb(pd.DataFrame(rows), "Categorical and binary variables")
    if "Missing values" in inc:
        h1("Missing values")
        fm = E.missing_fig(df)
        if fm is None:
            p("No missing values.")
        else:
            im(png(fm), "Percentage of missing values per column"); p("Columns with many missing values may need dropping or imputation.")
    if "Normality assessment" in inc:
        h1("Normality assessment")
        if nt.empty:
            p("No numeric variables to assess.")
        else:
            share = overall(nt, 80)[1]
            p(f"**{share:.0f}%** of the assessed numeric variables look normal-like. Rough screen only (80% threshold): "
              f"**{'mostly normal-like, so parametric tests are a reasonable starting point' if normal else 'mostly not normal-like, so non-parametric tests are the safer starting point'}**.")
            tb(nt.drop(columns=["rule"]), "Normality by variable")
            p("Up to 200 values the Shapiro-Wilk p-value decides; with more data, skewness and kurtosis are used (rules of thumb). "
              "Normality really matters per group or for model residuals, not for the whole dataset.")
    if "Distributions (first variables)" in inc and nc:
        h1("Distributions")
        for c in nc[:cfg.get("max_vars", 6)]:
            if df[c].notna().sum() >= 8:
                B.append(("h2", c)); im(stitch([png(f) for f in E.dist_figures(df[c], c)]), f"Histogram, boxplot and Q-Q plot: {c}")
                bl(E.interpret_numeric(c, E.summarize(df[c])))
    if "Correlation" in inc and len(nc) >= 2:
        h1("Correlation")
        f, cm = E.corr_heatmap(df, nc[:20], "pearson"); im(png(f), "Pearson correlation heatmap")
        tp = E.top_pairs(cm); tb(tp, "Strongest relationships"); bl(E.interpret_pairs(tp))
    if "Test advisor" in inc:
        h1("Suggested statistical tests")
        p(f"Based on the rough normality screen ({'mostly normal-like' if normal else 'mostly not normal-like'}), these tests are suggested. Treat them as guidance: check normality per group or in the residuals before deciding.")
        tb(pd.DataFrame([{"goal": g, "route": recommend(g, normal)["family"], "suggested test": recommend(g, normal)["test"],
                          "follow-up": recommend(g, normal)["post"]} for g in GOALS]), "Test suggestions")
    if "PCA" in inc and len(nc) >= 3:
        h1("Principal component analysis")
        try:
            res = PC.run_pca(df, nc[:30], True, bool(cfg.get("pca_log"))); m = PC.choose_m(res, "Cumulative variance threshold", 80)
            suit = PC.suitability(res); imp = PC.importance(res, m)
            im(png(PC.scree_fig(res, m)), "Scree plot"); im(png(PC.importance_fig(imp)), "Variable importance")
            tb(pd.DataFrame({"component": [f"PC{i + 1}" for i in range(res["p"])], "eigenvalue": res["eig"].round(3),
                             "% variance": (res["ratio"] * 100).round(2), "cumulative %": (res["cum"] * 100).round(2)}).head(10), "Variance explained")
            tb(imp.head(15), "Most important variables"); bl(PC.interpret_pca(res, m, imp, suit))
        except ValueError as e:
            p(f"PCA could not be run: {e}")
    if cfg.get("reg"):
        h1("Regression"); y, xs = cfg["reg"]; cats = cfg.get("reg_cat") or []
        try:
            dfr, dums = RG.add_dummies(df, cats) if cats else (df, [])
            allx = list(xs) + dums
            r = RG.fit_ols(dfr, y, allx); r["dummies"] = dums
            p(f"Response: **{y}**; numeric predictors: **{', '.join(xs) if xs else 'none'}**" + (f"; categorical predictors (dummy-coded, first level = reference): **{', '.join(cats)}**" if cats else "")
              + f"; rows used: {r['n']}; R-squared = {r['r2']:.3f}; adjusted = {r['adj_r2']:.3f}.")
            im(png(RG.line_fig(r) if (r["k"] == 1 and not dums) else RG.avp_fig(r)), "Regression line" if (r["k"] == 1 and not dums) else "Actual vs predicted")
            tb(r["coef"].round(4), "Coefficients"); im(stitch([png(f) for f in RG.diag_figs(r)]), "Residual diagnostics")
            bl(RG.interpret(r, show_dw=bool(cfg.get("reg_dw"))))
            if cfg.get("reg_ho"):
                try:
                    ho = RG.holdout_eval(dfr, y, allx, cfg.get("reg_frac", .25), cfg.get("reg_seed", 42))
                    tb(pd.DataFrame({"measure": ["Rows in training / test", "RMSE", "R-squared"],
                                     "training": [ho["n_train"], round(ho["rmse_train"], 4), round(ho["r2_train"], 3)],
                                     "test (unseen)": [ho["n_test"], round(ho["rmse_test"], 4), round(ho["r2_test"], 3)]}), "Hold-out test")
                    bl(RG.interpret_holdout(ho))
                except ValueError as e:
                    p(f"Hold-out test could not be run: {e}")
        except ValueError as e:
            p(f"Regression could not be run: {e}")
    if cfg.get("paired"):
        h1("Paired test"); a, b2 = cfg["paired"]
        try:
            r = paired_test(df[a].to_numpy(float), df[b2].to_numpy(float))
            p(f"Comparing **{a}** and **{b2}** on {r['n']} complete pairs (mean {a} = {r['mean1']:.4g}, mean {b2} = {r['mean2']:.4g}).")
            tb(pd.DataFrame({"result": ["Paired t-test p-value", "Mean difference (first - second)", "95% CI of the difference", "Cohen's dz",
                                        "Wilcoxon signed-rank p-value", "Matched-pairs rank-biserial", "Shapiro-Wilk p of the differences"],
                             "value": [fmt_p(r["p_t"]), f"{r['mean_diff']:.4g}", f"({r['ci'][0]:.4g}, {r['ci'][1]:.4g})", f"{r['cohen_dz']:.3f}",
                                       "-" if np.isnan(r["p_wilcoxon"]) else fmt_p(r["p_wilcoxon"]), "-" if np.isnan(r["rank_biserial"]) else f"{r['rank_biserial']:.3f}",
                                       "-" if np.isnan(r["shapiro_p_diff"]) else fmt_p(r["shapiro_p_diff"])]}), "Paired test")
            bl(interpret_paired(r, a, b2))
        except ValueError as e:
            p(f"Paired test could not be run: {e}")
    if cfg.get("one"):
        h1("One-sample test"); c, mu0 = cfg["one"]
        try:
            r = one_sample_test(df[c].to_numpy(float), float(mu0))
            p(f"Testing whether the mean of **{c}** differs from **{mu0:g}** (n = {r['n']}, mean = {r['mean']:.4g}, median = {r['median']:.4g}).")
            tb(pd.DataFrame({"result": ["One-sample t-test p-value", "95% CI of the mean", "Cohen's d", "Wilcoxon signed-rank p-value", "Shapiro-Wilk p"],
                             "value": [fmt_p(r["p_t"]), f"({r['ci'][0]:.4g}, {r['ci'][1]:.4g})", f"{r['cohen_d']:.3f}",
                                       "-" if np.isnan(r["p_wilcoxon"]) else fmt_p(r["p_wilcoxon"]), "-" if np.isnan(r["shapiro_p"]) else fmt_p(r["shapiro_p"])]}), "One-sample test")
            bl(interpret_one_sample(r))
        except ValueError as e:
            p(f"One-sample test could not be run: {e}")
    if cfg.get("friedman"):
        h1("Friedman test (repeated measures)"); fc = cfg["friedman"]
        try:
            r = friedman_test(df[fc])
            p(f"{r['n']} complete subjects measured under {r['k']} conditions: {', '.join(fc)}.")
            tb(r["summary"], "Summary by condition"); tb(r["pairs"], "Pairwise Wilcoxon signed-rank tests (Holm-adjusted)"); bl(interpret_friedman(r))
        except ValueError as e:
            p(f"Friedman test could not be run: {e}")
    if cfg.get("grp"):
        h1("Group comparison"); gn, gc = cfg["grp"]; d = df[[gn, gc]].dropna(); lv = sorted(d[gc].unique(), key=str)
        if len(lv) == 2:
            r = two_group_test(d.loc[d[gc] == lv[0], gn], d.loc[d[gc] == lv[1], gn]); lo, hi = r["mean_diff_CI"]
            p(f"Comparing **{gn}** between **{lv[0]}** (n={r['n1']}, mean={r['mean1']:.3f}) and **{lv[1]}** (n={r['n2']}, mean={r['mean2']:.3f}).")
            tb(pd.DataFrame({"result": ["Welch t-test p-value", "95% CI of mean difference", "Cohen's d", "Mann-Whitney p-value", "Rank-biserial correlation",
                                        "Shapiro-Wilk p (group 1 / group 2)", "Levene p"],
                             "value": [fmt_p(r['welch_p']), f"({lo:.3f}, {hi:.3f})", f"{r['cohen_d']:.3f}", fmt_p(r['mannwhitney_p']),
                                       f"{r['rank_biserial']:.3f}", f"{r['shapiro_p_group1']:.3f} / {r['shapiro_p_group2']:.3f}", f"{r['levene_p']:.3f}"]}), "Two-group test")
        elif 3 <= len(lv) <= 10:
            r = multi_group_test({str(k): s[gn].values for k, s in d.groupby(gc)})
            tb(r["groups"], "Group summary"); tb(r["tests"], "Overall tests"); tb(r["pairs"], "Pairwise comparisons (Holm-adjusted)"); bl(interpret_anova(r))
        else:
            p("The grouping column needs 2 to 10 groups.")
        if 2 <= len(lv) <= 10:
            gt = group_normality(df, gn, gc); tb(gt.drop(columns=["rule"]), f"Normality inside each group of {gc} (what t-tests and ANOVA need)")
            im(png(E.group_box(df, gn, gc)[0]), f"{gn} by {gc}")
    if cfg.get("chi"):
        h1("Chi-square test"); a, b = cfg["chi"]; d = df[[a, b]].dropna(); r = chi_square_test(d[a], d[b])
        tb(r["table"], "Observed counts"); tb(r["residuals"], "Adjusted standardised residuals"); bl(interpret_chi(r))
    B.append(("h1", "Notes"))
    B.append(("bullets", ["Interpretations are generated automatically from rules of thumb; check the plots and use your own judgement.",
                          "Association is not causation. Results describe this dataset only.",
                          "Regression uses ordinary least squares; categorical predictors are dummy-coded against the first level. A hold-out test shows accuracy on unseen rows.",
                          "Generated by StatGuide."]))
    return B


# ------------------------------------------------------------------ renderers
def _runs(par, text):
    for i, part in enumerate(text.split("**")):
        if part:
            par.add_run(part).bold = (i % 2 == 1)


def to_docx(blocks):
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    for b in blocks:
        t = b[0]
        if t == "title":
            doc.add_heading(b[1], 0); doc.add_paragraph(b[2])
        elif t in ("h1", "h2"):
            doc.add_heading(b[1], 1 if t == "h1" else 2)
        elif t == "p":
            _runs(doc.add_paragraph(), b[1])
        elif t == "bullets":
            for s in b[1]:
                _runs(doc.add_paragraph(style="List Bullet"), s)
        elif t == "table":
            cols, rows, note = fmt(b[1]); doc.add_paragraph().add_run(b[2]).bold = True
            tbl = doc.add_table(rows=1, cols=len(cols)); tbl.style = "Table Grid"
            for j, c in enumerate(cols):
                cell = tbl.rows[0].cells[j]; cell.text = ""; run = cell.paragraphs[0].add_run(c); run.bold = True; run.font.size = Pt(8)
            for row in rows:
                cells = tbl.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = ""; cells[j].paragraphs[0].add_run(v).font.size = Pt(8)
            if note:
                doc.add_paragraph(note)
            doc.add_paragraph()
        elif t == "img":
            doc.add_picture(io.BytesIO(b[1]), width=Inches(6.3)); doc.add_paragraph().add_run(b[2]).italic = True
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def to_pdf(blocks):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    ss = getSampleStyleSheet(); small = ParagraphStyle("small", parent=ss["Normal"], fontSize=6.5, leading=8)
    cap = ParagraphStyle("cap", parent=ss["Italic"], fontSize=8)

    def x(s):
        s = s.encode("latin-1", "replace").decode("latin-1").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return "".join(f"<b>{part}</b>" if i % 2 else part for i, part in enumerate(s.split("**")))
    story, W, hold = [], 17 * cm, None
    for b in blocks:
        t = b[0]
        if t == "title":
            story += [Paragraph(x(b[1]), ss["Title"]), Paragraph(x(b[2]), ss["Normal"]), Spacer(1, 10)]
        elif t == "h1":
            story += [Spacer(1, 8), Paragraph(x(b[1]), ss["Heading1"])]
        elif t == "h2":
            hold = Paragraph(x(b[1]), ss["Heading2"])
        elif t == "p":
            story += [Paragraph(x(b[1]), ss["Normal"]), Spacer(1, 4)]
        elif t == "bullets":
            story += [Paragraph(x(s), ss["Normal"], bulletText="\u2022") for s in b[1]] + [Spacer(1, 4)]
        elif t == "table":
            cols, rows, note = fmt(b[1]); lens = [max([len(c)] + [len(r[j]) for r in rows]) for j, c in enumerate(cols)]
            w = [min(max(n, 8), 28) for n in lens]; w = [W * v / sum(w) for v in w]
            data = [[Paragraph(f"<b>{x(c)}</b>", small) for c in cols]] + [[Paragraph(x(v), small) for v in r] for r in rows]
            tt = Table(data, colWidths=w, repeatRows=1)
            tt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .25, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [Paragraph(f"<b>{x(b[2])}</b>", ss["Normal"]), Spacer(1, 3), tt] + ([Paragraph(x(note), cap)] if note else []) + [Spacer(1, 8)]
        elif t == "img":
            iw, ih = PILImage.open(io.BytesIO(b[1])).size; wp = min(W, iw * 72 / 110); hp = wp * ih / iw
            if hp > 12 * cm:
                hp = 12 * cm; wp = hp * iw / ih
            story += [KeepTogether(([hold] if hold else []) + [Image(io.BytesIO(b[1]), width=wp, height=hp), Paragraph(x(b[2]), cap)]), Spacer(1, 8)]; hold = None
    out = io.BytesIO()
    def pg(c, d):
        c.setFont("Helvetica", 8); c.drawRightString(A4[0] - 2 * cm, 1 * cm, f"Page {d.page}")
    SimpleDocTemplate(out, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm).build(story, onFirstPage=pg, onLaterPages=pg)
    return out.getvalue()


def to_xlsx(blocks):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Report"; ws.sheet_view.showGridLines = False; NC = 10
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 16
    r = 1
    def text(s, bold=False, size=11, bullet=False):
        nonlocal r
        s = ("\u2022 " if bullet else "") + s.replace("**", ""); c = ws.cell(r, 1, s); c.font = Font(bold=bold, size=size)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(s) / 130)) * (size / 11); r += 1
    for b in blocks:
        t = b[0]
        if t == "title":
            text(b[1], True, 18); text(b[2]); r += 1
        elif t == "h1":
            r += 1; text(b[1], True, 14)
        elif t == "h2":
            text(b[1], True, 12)
        elif t == "p":
            text(b[1])
        elif t == "bullets":
            for s in b[1]:
                text(s, bullet=True)
        elif t == "table":
            cols, rows, note = fmt(b[1]); text(b[2], True)
            for j, c in enumerate(cols, 1):
                cell = ws.cell(r, j, c); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9D9D9")
            r += 1
            for row in rows:
                for j, v in enumerate(row, 1):
                    try:
                        v = float(v) if v not in ("", "nan") and v.replace(".", "", 1).replace("-", "", 1).replace("e", "", 1).replace("+", "", 1).isdigit() else v
                    except ValueError:
                        pass
                    ws.cell(r, j, v)
                r += 1
            if note:
                text(note)
            r += 1
        elif t == "img":
            img = XLImage(io.BytesIO(b[1])); sc = min(1, 850 / img.width); img.width, img.height = int(img.width * sc), int(img.height * sc)
            ws.add_image(img, f"A{r}"); r += math.ceil(img.height / 20) + 1; text(b[2]); r += 1
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
